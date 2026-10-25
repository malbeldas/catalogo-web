"""Genera plantilla/planilla-comercio.xlsx (SPEC.md §3 y BR-PRE-001).

Uso: python3 plantilla/generar.py [filas_web] [salida]
Requiere openpyxl. Solo hace falta para volver a generar la plantilla; la web no lo usa.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

FILAS_WEB = int(sys.argv[1]) if len(sys.argv) > 1 else 3000
SALIDA = sys.argv[2] if len(sys.argv) > 2 else 'plantilla/planilla-comercio.xlsx'

COLUMNAS = ['Código', 'Nombre', 'Categoría', 'Descripción', 'Precio', 'Precio anterior', 'Foto', 'Visible']
LETRA = dict(zip(COLUMNAS, 'ABCDEFGH'))
EJEMPLOS = [
    ['A001', 'Remera de algodón lisa', 'Ropa', 'Talles S a XL', 12900, 15900, '', 'SÍ'],
    ['A002', 'Taza de cerámica 350 ml', 'Bazar', 'Apta microondas', 6500, '', '', 'SÍ'],
    ['A003', 'Producto escondido de ejemplo', 'Bazar', 'Con Visible = NO no aparece en la web', 9900, '', '', 'NO'],
    ['', 'Velas aromáticas x3', 'Deco', 'Sin código: la web usa el número de fila', '', '', '', ''],
]
AJUSTES = [
    ['WhatsApp', '', 'Con código de país, sin + ni espacios. Ejemplo: 5491112345678'],
    ['Precios', 'SÍ', 'SÍ muestra los precios en la web. NO los esconde del todo, también de la pestaña Web.'],
    ['Nombre', '', 'Opcional: reemplaza el nombre del comercio en la web.'],
    ['Bajada', '', 'Opcional: la frase que va debajo del nombre.'],
]
AYUDA = [
    'Cómo usar esta planilla',
    '',
    '1. Cargá tus productos en la pestaña Productos, uno por fila. El Nombre es obligatorio: una fila sin nombre no se muestra.',
    '2. Precio: un número, con o sin $ y puntos (14900 o $ 14.900). Vacío o con texto, la web muestra "Consultar".',
    '3. Precio anterior es opcional: la web lo muestra tachado, con el porcentaje de descuento.',
    '4. Foto: el link de Google Drive compartido como "Cualquier persona con el enlace", o el link de una imagen.',
    '5. Visible: poné NO para esconder un producto sin borrarlo. Vacío cuenta como SÍ.',
    '6. Código: si se repite, la web muestra solo la primera fila con ese código.',
    '7. En Ajustes cargá tu WhatsApp y elegí si se ven los precios.',
    '8. No edites la pestaña Web: se arma sola desde Productos y Ajustes. Web y Ajustes son las únicas pestañas que se publican.',
    '9. Los cambios tardan hasta 5 minutos en verse en la web.',
    f'10. La pestaña Web está preparada para {FILAS_WEB:,} productos. Si tenés más, copiá su última fila hacia abajo.'.replace(',', '.'),
]

TITULO = Font(bold=True, color='FFFFFF')
FONDO = PatternFill('solid', fgColor='1F2937')


def titulos(ws, valores, anchos):
    ws.append(valores)
    for c, ancho in zip(ws[1], anchos):
        c.font, c.fill = TITULO, FONDO
        ws.column_dimensions[c.column_letter].width = ancho
    ws.freeze_panes = 'A2'


def lista_si_no(ws, rango):
    dv = DataValidation(type='list', formula1='"SÍ,NO"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rango)


wb = Workbook()
prod = wb.active
prod.title = 'Productos'
titulos(prod, COLUMNAS, [10, 36, 16, 36, 12, 16, 40, 10])
for fila in EJEMPLOS:
    prod.append(fila)
for col in 'EF':
    for c in prod[col][1:]:
        c.number_format = '"$"#,##0'
lista_si_no(prod, f'H2:H{FILAS_WEB + 1}')

aj = wb.create_sheet('Ajustes')
titulos(aj, ['Ajuste', 'Valor', 'Para qué sirve'], [14, 22, 80])
for fila in AJUSTES:
    aj.append(fila)
lista_si_no(aj, 'B3')

# Web: copia de Productos armada con fórmulas, que es lo único que se publica (BR-PRE-001).
# - Cada fila lee la misma fila de Productos con INDEX(...;ROW()): insertar o borrar filas en Productos no la desarma.
# - Fila sin nombre o con Visible = NO: queda vacía entera.
# - Precio y Precio anterior se copian solo si Ajustes > Precios es SÍ; si no, ni siquiera llegan al CSV publicado.
# - &"" deja todo como texto: una celda vacía de Productos sale vacía y no como 0.
web = wb.create_sheet('Web')
web.sheet_properties.tabColor = '9CA3AF'
publicadas = COLUMNAS[:-1]
titulos(web, publicadas, [10, 36, 16, 36, 12, 16, 40])
celda = lambda col: f'INDEX(Productos!${LETRA[col]}:${LETRA[col]},ROW())'
visible = f'AND({celda("Nombre")}<>"",UPPER(TRIM({celda("Visible")}))<>"NO")'
precios = 'OR(UPPER(TRIM(IFERROR(VLOOKUP("Precios",Ajustes!$A:$B,2,FALSE),"")))="SÍ",UPPER(TRIM(IFERROR(VLOOKUP("Precios",Ajustes!$A:$B,2,FALSE),"")))="SI")'
for r in range(2, FILAS_WEB + 2):
    for i, col in enumerate(publicadas, start=1):
        cond = f'AND({visible},{precios})' if col in ('Precio', 'Precio anterior') else visible
        web.cell(row=r, column=i, value=f'=IF({cond},{celda(col)}&"","")')

ay = wb.create_sheet('Ayuda')
ay.column_dimensions['A'].width = 120
for linea in AYUDA:
    ay.append([linea])
ay['A1'].font = Font(bold=True, size=14)
for c in ay['A']:
    c.alignment = Alignment(wrap_text=True, vertical='top')

wb.save(SALIDA)
print(SALIDA, FILAS_WEB, 'filas en Web')
